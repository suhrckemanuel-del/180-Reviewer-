# PROTOTYPE test fixture for build_deck_walk.py. Writes an INVENTED 18-slide deck (the food bank volunteer-app
# deck the issue 23 drafts review) as test-deck.pptx (python-pptx) and test-deck.pdf (PyMuPDF), with the same
# words on the same slides, so the review's quotes are found on their cited slides. No client material.
#   python make_test_deck.py
import fitz
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

SLIDES = [
    ["Volunteer scheduling at Oakfield Food Bank", "180DC · Deliverable 2"],
    ["Executive summary", "Shift coverage is slipping as volunteer numbers grow", "A three-month pilot will test the app with two shift teams.", "We recommend a volunteer booking app"],
    ["Project scope", "How can coordinators fill shifts with less effort?"],
    ["Approach", "6 coordinator interviews", "12 volunteer survey responses", "3 scheduling tools compared"],
    ["Section 1 · Situation"],
    ["What coordinators told us", "Coordinators spend about 4 hours a week chasing replies", "Last-minute gaps are filled by the same few people"],
    ["Shift coverage", "30% of shifts are unfilled at least once a month", "Gaps cluster on weekday mornings"],
    ["Volunteer demographics", "Age bands", "Years volunteering"],
    ["Volunteer attitudes", "Most volunteers would use an app (9 of 12 surveyed)", "Reminders were the most wanted feature"],
    ["Current tools", "WhatsApp group", "Paper rota on the warehouse wall"],
    ["Weekly demand against volunteers"],                      # chart only, drawn as shapes
    ["Option comparison: WhatsApp, shared spreadsheet, booking app", "Cost · effort · reach"],
    ["Section 3 · Recommendation"],
    ["Recommendation", "We recommend adopting a volunteer booking app.", "Automated reminders", "Self-service shift swaps"],
    ["Next steps", "Next steps: explore options with the board"],
    ["Thank you", "Questions?"],
    ["Appendix A · Survey method", "Survey sent to the volunteer WhatsApp group", "12 responses over two weeks"],
    ["Appendix B · Interview guide"],
]
BARS = [0.55, 0.8, 0.7, 0.9, 0.6, 0.35, 0.3]
GREEN = (0x78 / 255, 0xB0 / 255, 0x38 / 255)


def pdf(path):
    doc = fitz.open()
    for i, (title, *items) in enumerate(SLIDES, 1):
        p = doc.new_page(width=960, height=540)
        p.draw_rect(fitz.Rect(0, 0, 14, 540), color=None, fill=GREEN)
        # titles stay on one line: a wrapped title extracts as two lines and the quote check misses it (issue 20)
        p.insert_textbox(fitz.Rect(60, 50, 940, 140), title, fontsize=30 if len(title) < 48 else 24, fontname='hebo')
        y = 160
        for it in items:
            p.insert_textbox(fitz.Rect(80, y, 900, y + 60), '- ' + it, fontsize=19, fontname='helv')
            y += 52
        if i == 11:
            for k, h in enumerate(BARS):
                p.draw_rect(fitz.Rect(120 + k * 100, 470 - h * 290, 190 + k * 100, 470), color=None, fill=(0.25, 0.27, 0.24))
        p.insert_text(fitz.Point(900, 515), str(i), fontsize=11, fontname='helv', color=(0.4, 0.4, 0.4))
    doc.save(path)


def pptx(path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for i, (title, *items) in enumerate(SLIDES, 1):
        s = prs.slides.add_slide(prs.slide_layouts[5])          # title only
        s.shapes.title.text = title
        if items:
            tf = s.shapes.add_textbox(Inches(0.9), Inches(2.1), Inches(11.5), Inches(4.5)).text_frame
            tf.word_wrap = True
            for k, it in enumerate(items):
                para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                para.text, para.font.size = it, Pt(24)
        if i == 11:
            for k, h in enumerate(BARS):
                bar = s.shapes.add_shape(1, Inches(1.6 + k * 1.4), Inches(6.6 - h * 4), Inches(0.9), Inches(h * 4))
                bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor(0x40, 0x45, 0x3D); bar.line.fill.background()
    prs.save(path)


if __name__ == '__main__':
    pdf('test-deck.pdf')
    pptx('test-deck.pptx')
    print(f'wrote test-deck.pdf and test-deck.pptx ({len(SLIDES)} slides, invented)')
