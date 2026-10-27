"""Profile a data file before computing anything from it.

Usage: python profile_data.py <file.csv | file.xlsx> [sheet name]
Prints rows, columns, each column's type, gaps and range, and the first rows.
CSV needs only the standard library; .xlsx needs openpyxl (it says so if it's missing).
"""
import csv, io, sys


def read(path, sheet=None):
    if path.lower().endswith((".xlsx", ".xlsm")):
        try:
            import openpyxl
        except ImportError:
            sys.exit("openpyxl is not installed: save the sheet as CSV, or paste the table.")
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        if len(wb.sheetnames) > 1 and not sheet:
            print("Sheets:", ", ".join(wb.sheetnames), "(profiling the first; the others' first rows follow)")
            for other in wb.sheetnames[1:]:
                for r in list(wb[other].iter_rows(values_only=True))[:3]:
                    print(f"  [{other}]", " | ".join("" if c is None else str(c) for c in r))
        ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
        print("Reading sheet:", ws.title)
        rows = [["" if c is None else c for c in r] for r in ws.iter_rows(values_only=True)]
    else:
        text = open(path, encoding="utf-8-sig", newline="").read()
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    return rows[0], rows[1:]


def as_number(v):
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "").replace("%", "")
    if s.count(",") and s.count(".") == 0 and len(s.split(",")[-1]) != 3:
        s = s.replace(",", ".")  # decimal comma
    s = s.replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    head, body = read(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
    print(f"Rows: {len(body)}  Columns: {len(head)}")
    for i, name in enumerate(head):
        col = [r[i] if i < len(r) else "" for r in body]
        filled = [c for c in col if str(c).strip() != ""]
        nums = [n for n in (as_number(c) for c in filled) if n is not None]
        gaps = len(col) - len(filled)
        if filled and len(nums) == len(filled):
            print(f"- {name}: number, min {min(nums):g}, max {max(nums):g}, sum {sum(nums):g}, gaps {gaps}")
        else:
            distinct = sorted({str(c) for c in filled})
            shown = ", ".join(distinct[:10]) + (" ..." if len(distinct) > 10 else "")
            print(f"- {name}: text, {len(distinct)} distinct ({shown}), gaps {gaps}")
    print("First rows:")
    for r in body[:5]:
        print("  ", " | ".join(str(c) for c in r))


main()
