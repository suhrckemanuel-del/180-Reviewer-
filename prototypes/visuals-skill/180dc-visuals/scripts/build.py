"""Build the slides: check every number, then write chart PNGs and a .pptx with native charts.

Usage: python build.py slides.json numbers.json <out dir>
Stops before writing anything if a number on a slide isn't in numbers.json.
PNGs need matplotlib; the .pptx needs python-pptx. A missing library skips that output and
says so; the other one is still written.
"""
import json, os, sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from check_numbers import check

GREEN, GREY, DARK, INK, MUTED, RULE = "78B038", "A3A8AE", "3A3F45", "1D2124", "6B7075", "E3E5E1"
PLACEHOLDER = "[Your title: what does this chart prove?]"
TYPES = ("line", "column", "bar", "stacked100")


def fail(msg):
    sys.exit("STOPPED: " + msg)


def fmt(v, unit, dec):
    s = f"{v:,.{dec}f}"
    if unit in ("€", "$", "£"):
        return unit + s
    return s + unit if unit in ("%", "") else f"{s} {unit}"


def resolve(slide, nums, n):
    """Validate one slide and return its chart data: categories, series [(name, values, highlight)]."""
    c = slide.get("chart")
    if not c:
        return None
    if c.get("type") not in TYPES:
        fail(f"slide {n}: chart type must be one of {', '.join(TYPES)}")
    cats = nums.get(c.get("categories"))
    if not isinstance(cats, list):
        fail(f"slide {n}: categories key {c.get('categories')!r} is not a list in numbers.json")
    series = []
    for s in c.get("series", []):
        vals = nums.get(s.get("key"))
        if not isinstance(vals, list) or len(vals) != len(cats):
            fail(f"slide {n}: series key {s.get('key')!r} must be a list as long as the categories")
        series.append((s["name"], [float(v) for v in vals], bool(s.get("highlight"))))
    if not series:
        fail(f"slide {n}: the chart has no series")
    if c["type"] == "stacked100":
        for i, cat in enumerate(cats):
            total = sum(s[1][i] for s in series)
            if abs(total - 100) > 0.6:
                fail(f"slide {n}: parts of {cat!r} add to {total:.1f}, not 100")
    ev = c.get("event")
    if ev and ev.get("at") not in cats:
        fail(f"slide {n}: event 'at' {ev.get('at')!r} is not one of the categories")
    return [str(x) for x in cats], series


def slide_texts(slide):
    c = slide.get("chart") or {}
    out = [slide.get("title") or "", slide.get("caption") or "", slide.get("source") or "",
           c.get("y_label") or "", (c.get("event") or {}).get("label") or ""]
    out += [s.get("name", "") for s in c.get("series", [])]
    return [t for t in out if t]


def colour(highlight, any_highlight, i):
    if highlight:
        return GREEN
    return GREY if any_highlight else (GREY, DARK, "C9CDD1")[min(i, 2)]


# ---------- PNG (matplotlib) ----------

def build_png(slide, data, path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.ticker import FuncFormatter
    from matplotlib import font_manager
    have = {f.name for f in font_manager.fontManager.ttflist}
    family = next((f for f in ("IBM Plex Sans", "Arial", "Liberation Sans") if f in have), "DejaVu Sans")
    plt.rcParams.update({"font.family": family,
                         "font.size": 15, "axes.edgecolor": "#" + RULE, "text.color": "#" + INK})
    c = slide["chart"]
    cats, series = data
    unit, dec = c.get("unit", ""), int(c.get("decimals", 0))
    anyh = any(h for _, _, h in series)
    n = len(cats)
    fig, ax = plt.subplots(figsize=(12, 5.6), dpi=200)
    for side in ("top", "right", "left"):
        ax.spines[side].set_visible(False)
    ax.tick_params(colors="#" + MUTED, length=0)
    ax.grid(axis="x" if c["type"] in ("bar", "stacked100") else "y", color="#" + RULE, linewidth=1)
    ax.set_axisbelow(True)
    tick = FuncFormatter(lambda v, _: fmt(v, unit, 0))
    t = c["type"]
    if t == "line":
        x = list(range(n))
        ends = []
        for i, (name, vals, h) in enumerate(series):
            ax.plot(x, vals, color="#" + colour(h, anyh, i), linewidth=3.5 if h else 2.2, zorder=3 if h else 2)
            ax.plot([x[-1]], [vals[-1]], "o", color="#" + colour(h, anyh, i), markersize=7, zorder=4)
            ends.append([vals[-1], f"{name}  {fmt(vals[-1], unit, dec)}", h])
        top = c.get("y_max") or max(max(v) for _, v, _ in series) * 1.15
        ax.set_ylim(0, top)
        gap = top * 0.06  # keep end labels apart
        ends.sort(key=lambda e: e[0])
        for k in range(1, len(ends)):
            ends[k][0] = max(ends[k][0], ends[k - 1][0] + gap)
        for y, label, h in ends:
            ax.text(n - 1 + 0.25, y, label, va="center", fontsize=15, fontweight="bold" if h else "normal")
        ax.set_xlim(-0.3, n - 1 + max(1.6, n * 0.16))
        ax.set_xticks(x, cats)
        ax.yaxis.set_major_formatter(tick)
    elif t == "column":
        w = 0.8 / len(series)
        for i, (name, vals, h) in enumerate(series):
            xs = [k - 0.4 + w * (i + 0.5) for k in range(n)]
            bars = ax.bar(xs, vals, width=w * 0.92, color="#" + colour(h, anyh, i), label=name)
            ax.bar_label(bars, labels=[fmt(v, unit, dec) for v in vals], padding=4, fontsize=14)
        ax.set_xticks(range(n), cats)
        ax.set_ylim(0, c.get("y_max") or max(max(v) for _, v, _ in series) * 1.18)
        ax.yaxis.set_major_formatter(tick)
        if len(series) > 1:
            ax.legend(frameon=False, loc="upper left", ncol=len(series), bbox_to_anchor=(0, 1.1))
    elif t == "bar":
        hc = c.get("highlight_category")
        w = 0.8 / len(series)
        for i, (name, vals, h) in enumerate(series):
            ys = [k - 0.4 + w * (i + 0.5) for k in range(n)]
            cols = ["#" + (GREEN if hc and cats[k] == hc else GREY if hc else colour(h, anyh, i)) for k in range(n)]
            bars = ax.barh(ys, vals, height=w * 0.9, color=cols, label=name)
            ax.bar_label(bars, labels=[fmt(v, unit, dec) for v in vals], padding=6, fontsize=14)
        ax.set_yticks(range(n), cats)
        ax.invert_yaxis()
        ax.set_xlim(0, c.get("y_max") or max(max(v) for _, v, _ in series) * 1.15)
        ax.xaxis.set_major_formatter(tick)
        ax.spines["bottom"].set_visible(False)
        if len(series) > 1:
            ax.legend(frameon=False, loc="lower right")
    else:  # stacked100
        left = [0.0] * n
        for i, (name, vals, h) in enumerate(series):
            col = colour(h, anyh, i) if anyh else (DARK, GREEN, GREY, "C9CDD1")[min(i, 3)]
            bars = ax.barh(range(n), vals, left=left, height=0.6, color="#" + col, label=name,
                           edgecolor="white", linewidth=2)
            dark = col in (DARK,)
            ax.bar_label(bars, labels=[fmt(v, unit, dec) if v >= 6 else "" for v in vals],
                         label_type="center", fontsize=16, fontweight="bold", color="white" if dark else "#" + INK)
            left = [a + b for a, b in zip(left, vals)]
        ax.set_yticks(range(n), cats)
        ax.invert_yaxis()
        ax.set_xlim(0, 100)
        ax.xaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:.0f}%"))
        ax.spines["bottom"].set_visible(False)
        ax.legend(frameon=False, loc="lower left", ncol=len(series), bbox_to_anchor=(0, 1.0))
    ev = c.get("event")
    if ev and t in ("line", "column"):
        xe = cats.index(ev["at"]) - 0.5
        ax.axvline(xe, color="#" + INK, linestyle=(0, (4, 4)), linewidth=1.5, zorder=1)
        ax.text(xe + 0.08, ax.get_ylim()[1], ev["label"], va="top", fontsize=14, fontweight="bold")
    if c.get("y_label"):
        (ax.set_xlabel if t in ("bar", "stacked100") else ax.set_ylabel)(c["y_label"], color="#" + MUTED, fontsize=14)
    foot = "   ".join(x for x in (slide.get("caption"), slide.get("source")) if x)
    if foot:
        import textwrap
        foot = "\n".join(textwrap.wrap(foot, 150))
        fig.text(0.01, 0.005, foot, fontsize=11, color="#" + MUTED, ha="left", va="bottom")
    fig.tight_layout(rect=(0, 0.05 + 0.03 * (foot or "").count("\n"), 1, 1))
    fig.savefig(path, facecolor="white")
    plt.close(fig)


# ---------- PPTX (python-pptx, native charts) ----------

def build_pptx(slides, datas, path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_MARKER_STYLE
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Inches, Pt
    from lxml import etree

    rgb = RGBColor.from_string
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    CX, CY, CW, CH = Inches(0.6), Inches(1.95), Inches(12.1), Inches(4.5)
    PX, PY, PW, PH = 0.08, 0.1, 0.78, 0.76  # plot area inside the chart frame, as fractions

    def text(sl, x, y, w, h, s, size, color=INK, bold=False):
        tf = sl.shapes.add_textbox(x, y, w, h).text_frame
        tf.word_wrap = True
        r = tf.paragraphs[0].add_run()
        r.text, r.font.size, r.font.bold, r.font.color.rgb = s, Pt(size), bold, rgb(color)

    def pin_plot_area(chart):
        pa = chart._chartSpace.find(qn("c:chart")).find(qn("c:plotArea"))
        old = pa.find(qn("c:layout"))
        if old is not None:
            pa.remove(old)
        lay = etree.SubElement(pa, qn("c:layout"))
        pa.remove(lay)
        pa.insert(0, lay)
        ml = etree.SubElement(lay, qn("c:manualLayout"))
        for tag, val in (("layoutTarget", "inner"), ("xMode", "edge"), ("yMode", "edge"),
                         ("x", PX), ("y", PY), ("w", PW), ("h", PH)):
            etree.SubElement(ml, qn("c:" + tag)).set("val", str(val))

    for n, (s, data) in enumerate(zip(slides, datas), 1):
        sl = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only: a real title, so outline view shows the ghost deck
        title = sl.shapes.title
        title.left, title.top, title.width, title.height = Inches(0.6), Inches(0.35), Inches(12.1), Inches(1.15)
        title.text = s.get("title") or PLACEHOLDER
        for p in title.text_frame.paragraphs:
            p.alignment = PP_ALIGN.LEFT
            for r in p.runs:
                r.font.size, r.font.bold = Pt(28), True
                r.font.color.rgb = rgb(INK if s.get("title") else "9A3412")
        if s.get("notes"):
            sl.notes_slide.notes_text_frame.text = s["notes"]
        if data is None:
            continue
        c = s["chart"]
        cats, series = data
        unit, dec = c.get("unit", ""), int(c.get("decimals", 0))
        anyh = any(h for _, _, h in series)
        nf = f'#,##0{"." + "0" * dec if dec else ""}' + ('"%"' if unit == "%" else "")
        cd = CategoryChartData(number_format=nf)
        cd.categories = cats
        for name, vals, _ in series:
            cd.add_series(name, vals)
        kind = {"line": XL_CHART_TYPE.LINE, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
                "bar": XL_CHART_TYPE.BAR_CLUSTERED, "stacked100": XL_CHART_TYPE.BAR_STACKED_100}[c["type"]]
        ch = sl.shapes.add_chart(kind, CX, CY, CW, CH, cd).chart
        ch.font.size = Pt(14)
        ch.font.color.rgb = rgb(MUTED)
        horizontal = c["type"] in ("bar", "stacked100")
        if horizontal:
            ch.category_axis.reverse_order = True
        va = ch.value_axis
        va.has_major_gridlines = c["type"] != "stacked100"
        if va.has_major_gridlines:
            va.major_gridlines.format.line.color.rgb = rgb(RULE)
        va.format.line.fill.background()
        ch.category_axis.format.line.color.rgb = rgb(RULE)
        ch.category_axis.tick_labels.font.color.rgb = rgb(INK)
        if c["type"] == "stacked100":
            va.visible = False  # the segments carry their values
        else:
            va.minimum_scale = 0
            if c.get("y_max"):
                va.maximum_scale = c["y_max"]
            va.tick_labels.number_format = '#,##0' + ('"%"' if unit == "%" else "")
            va.tick_labels.number_format_is_linked = False
        ch.has_legend = len(series) > 1 and c["type"] != "line"
        if ch.has_legend:
            ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.TOP, False
        plot = ch.plots[0]
        if c["type"] in ("column", "bar"):
            plot.gap_width = 60
        if c["type"] == "stacked100":
            plot.gap_width = 50
            plot.overlap = 100
        hc = c.get("highlight_category")
        last_rank = sorted(range(len(series)), key=lambda i: series[i][1][-1])
        for i, ser in enumerate(plot.series):
            name, vals, h = series[i]
            col = colour(h, anyh, i) if (anyh or c["type"] != "stacked100") else (DARK, GREEN, GREY, "C9CDD1")[min(i, 3)]
            if c["type"] == "line":
                ser.format.line.color.rgb, ser.format.line.width = rgb(col), Pt(3.5 if h else 2.25)
                ser.smooth = False
                ser.marker.style = XL_MARKER_STYLE.NONE
                lab = ser.points[len(vals) - 1].data_label
                lab.text_frame.text = f"{name} {fmt(vals[-1], unit, dec)}"
                lab.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
                lab.text_frame.paragraphs[0].runs[0].font.bold = h
                lab.text_frame.paragraphs[0].runs[0].font.color.rgb = rgb(INK)
                rank = last_rank.index(i)
                lab.position = (XL_LABEL_POSITION.BELOW if rank == 0 else
                                XL_LABEL_POSITION.ABOVE if rank == len(series) - 1 else XL_LABEL_POSITION.RIGHT)
            else:
                ser.format.fill.solid()
                ser.format.fill.fore_color.rgb = rgb(col)
                if hc and c["type"] == "bar":
                    for k, cat in enumerate(cats):
                        pt = ser.points[k]
                        pt.format.fill.solid()
                        pt.format.fill.fore_color.rgb = rgb(GREEN if cat == hc else GREY)
                dl = ser.data_labels
                dl.show_value = True
                dl.number_format, dl.number_format_is_linked = nf, False
                dl.font.size = Pt(16 if c["type"] == "stacked100" else 14)
                dl.font.bold = c["type"] == "stacked100"
                dl.font.color.rgb = rgb("FFFFFF" if c["type"] == "stacked100" and col == DARK else INK)
                dl.position = XL_LABEL_POSITION.CENTER if c["type"] == "stacked100" else XL_LABEL_POSITION.OUTSIDE_END
        ev = c.get("event")
        if c["type"] in ("line", "column") or ev:
            pin_plot_area(ch)
        if ev and c["type"] in ("line", "column"):
            k = cats.index(ev["at"])
            x = CX + int(CW * (PX + PW * k / len(cats)))
            y0, y1 = CY + int(CH * PY), CY + int(CH * (PY + PH))
            ln = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x, y0, x, y1)
            ln.line.color.rgb, ln.line.width, ln.line.dash_style = rgb(INK), Pt(1.5), MSO_LINE_DASH_STYLE.DASH
            text(sl, x + Inches(0.08), y0 - Inches(0.05), Inches(4), Inches(0.4), ev["label"], 14, bold=True)
        if c.get("y_label"):
            text(sl, CX, CY - Inches(0.3), Inches(8), Inches(0.3), c["y_label"], 13, MUTED)
        foot_y = CY + CH + Inches(0.1)
        if s.get("caption"):
            text(sl, CX, foot_y, CW, Inches(0.35), s["caption"], 14)
            foot_y += Inches(0.38)
        if s.get("source"):
            text(sl, CX, foot_y, CW, Inches(0.35), s["source"], 11, MUTED)
    prs.save(path)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) < 4:
        sys.exit(__doc__)
    spec = json.load(open(sys.argv[1], encoding="utf-8"))
    nums = json.load(open(sys.argv[2], encoding="utf-8"))
    out = sys.argv[3]
    slides = spec.get("slides", [])
    if not slides:
        fail("slides.json has no slides")
    bad = []
    for n, s in enumerate(slides, 1):
        lines, ok = check(nums, slide_texts(s))
        if not ok:
            bad.append(f"slide {n}:\n" + "\n".join(l for l in lines if "MISSING" in l))
    if bad:
        fail("numbers on the slides that numbers.json doesn't hold. Compute them or take them out.\n" + "\n".join(bad))
    datas = [resolve(s, nums, n) for n, s in enumerate(slides, 1)]
    os.makedirs(out, exist_ok=True)
    made, skipped = [], []
    try:
        import matplotlib  # noqa: F401
        for n, (s, d) in enumerate(zip(slides, datas), 1):
            if d:
                p = os.path.join(out, f"chart-{n}.png")
                build_png(s, d, p)
                made.append(p)
    except ImportError:
        skipped.append("PNG charts (matplotlib is not installed)")
    try:
        import pptx  # noqa: F401
        p = os.path.join(out, "slides.pptx")
        build_pptx(slides, datas, p)
        made.append(p)
    except ImportError:
        skipped.append(".pptx (python-pptx is not installed)")
    print("ALL NUMBERS MATCH")
    print("Wrote: " + (", ".join(made) or "nothing"))
    if skipped:
        print("Skipped: " + "; ".join(skipped))
    untitled = [str(n) for n, s in enumerate(slides, 1) if not s.get("title")]
    if untitled:
        print("No title yet on slide " + ", ".join(untitled) + ": the placeholder is in.")


main()
