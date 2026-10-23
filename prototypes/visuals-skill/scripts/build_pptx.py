"""PROTOTYPE. The same four slides as a .pptx with native, editable charts (python-pptx).
This is the candidate Free-path output: it opens in PowerPoint and imports into Google Slides,
and the team can restyle the chart instead of pasting a picture.

Usage: python scripts/build_pptx.py   (run from the prototype folder)
"""
import json
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_MARKER_STYLE
from pptx.util import Emu, Pt

D = json.load(open("data/recomputed.json", encoding="utf-8"))
C = json.load(open("content/consultant.json", encoding="utf-8"))
S = {s["id"]: s for s in C["slides"]}
GREEN, PHONE, INK, MUTED, GREEN_TEXT = (RGBColor.from_string(x) for x in ("78B038", "3A3F45", "1D2124", "6B7075", "4A7322"))
FONT = "IBM Plex Sans"
SOURCE = (f"Source: pantry booking log, Jan–Dec, {D['online_booked_total']:,} online and "
          f"{D['phone_booked_total']:,} phone bookings. Invented data for a prototype.")

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
W = prs.slide_width
M = Emu(640000)

def text(slide, x, y, w, h, s, size, bold=False, color=INK, mono=False):
    tf = slide.shapes.add_textbox(x, y, w, h).text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    r.text = s
    r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, color
    r.font.name = "IBM Plex Mono" if mono else FONT
    return tf

def base(s, n):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    text(slide, M, Emu(380000), W - 2 * M, Emu(300000), s["tag"].upper(), 11, color=GREEN_TEXT, mono=True)
    text(slide, M, Emu(640000), W - 2 * M, Emu(900000), s["title"], 26, bold=True)
    slide.shapes.add_picture("assets/180dc-logo.png", M, Emu(6330000), height=Emu(260000))
    text(slide, W - M - Emu(6400000), Emu(6330000), Emu(6400000), Emu(300000),
         f"{C['org']} · invented data · prototype   {n}", 9, color=MUTED, mono=True)
    slide.notes_slide.notes_text_frame.text = s["notes"]
    return slide

def style_axis(ax, size=11):
    ax.tick_labels.font.size = Pt(size)
    ax.tick_labels.font.color.rgb = MUTED
    ax.format.line.color.rgb = RGBColor.from_string("D9DBD6")

# 1 answer
a = S["answer"]
sl = base(a, 1)
for k, r in enumerate(a["reasons"]):
    text(sl, M, Emu(1900000 + k * 900000), W - 2 * M, Emu(700000), f"0{k + 2}   {r}", 20)
text(sl, M, Emu(4800000), W - 2 * M, Emu(400000), a["footnote"], 12, color=MUTED)

# 2 line chart, native
sl = base(S["trend"], 2)
cd = CategoryChartData()
cd.categories = D["months"]
cd.add_series("Online", [v / 100 for v in D["online_rate"]])
cd.add_series("Phone", [v / 100 for v in D["phone_rate"]])
ch = sl.shapes.add_chart(XL_CHART_TYPE.LINE, M, Emu(1700000), W - 2 * M, Emu(4200000), cd).chart
ch.has_legend, ch.legend.position, ch.legend.include_in_layout = True, XL_LEGEND_POSITION.TOP, False
ch.legend.font.size = Pt(12)
va = ch.value_axis
va.maximum_scale, va.minimum_scale, va.major_unit = 0.25, 0, 0.05
va.tick_labels.number_format, va.tick_labels.number_format_is_linked = '0%', False
va.major_gridlines.format.line.color.rgb = RGBColor.from_string("E4E6E1")
style_axis(va); style_axis(ch.category_axis)
for ser, col in zip(ch.plots[0].series, (GREEN, PHONE)):
    ser.format.line.color.rgb, ser.format.line.width, ser.smooth = col, Pt(2.75), False
    ser.marker.style = XL_MARKER_STYLE.NONE
    last = ser.points[len(D["months"]) - 1].data_label
    last.text_frame.text = f'{ser.name} {D[ser.name.lower() + "_rate"][-1]}%'
    last.position = XL_LABEL_POSITION.BELOW if ser.name == 'Online' else XL_LABEL_POSITION.ABOVE
    last.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
text(sl, M, Emu(5900000), W - 2 * M, Emu(300000),
     "Text reminders started 1 July, online bookings only. Missed pickups as a share of booked pickups. " + SOURCE, 9, color=MUTED)

# 3 share bars, native 100% stacked
sl = base(S["share"], 3)
cd = CategoryChartData()
cd.categories = ["Missed pickups", "Bookings"]  # bar charts list bottom-up
p_b, p_m = D["phone_share_of_bookings_after"], D["phone_share_of_missed_after"]
cd.add_series("Phone", [p_m / 100, p_b / 100])
cd.add_series("Online", [1 - p_m / 100, 1 - p_b / 100])
ch = sl.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED_100, M, Emu(1800000), W - 2 * M, Emu(3900000), cd).chart
ch.has_legend, ch.legend.position, ch.legend.include_in_layout = True, XL_LEGEND_POSITION.TOP, False
ch.legend.font.size = Pt(12)
ch.plots[0].gap_width = 60
style_axis(ch.value_axis); style_axis(ch.category_axis, 14)
ch.value_axis.has_major_gridlines = False
for ser, col, fc in zip(ch.plots[0].series, (PHONE, GREEN), (RGBColor(255, 255, 255), INK)):
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = col
    dl = ser.data_labels
    dl.show_value, dl.number_format, dl.number_format_is_linked = True, '0.0%', False
    dl.font.size, dl.font.bold, dl.font.color.rgb = Pt(16), True, fc
text(sl, M, Emu(5900000), W - 2 * M, Emu(300000),
     f"Jul–Dec: {D['phone_missed_after']} of {D['missed_after_total']} missed pickups were phone bookings. " + SOURCE, 9, color=MUTED)

# 4 next step
sl = base(S["next"], 4)
rows = S["next"]["rows"]
tbl = sl.shapes.add_table(len(rows), 2, M, Emu(1900000), W - 2 * M, Emu(len(rows) * 560000)).table
tbl.columns[0].width = Emu(2700000)
tbl.columns[1].width = W - 2 * M - Emu(2700000)
for r, (k, v) in enumerate(rows):
    for c, val in enumerate((k.upper(), v)):
        cell = tbl.cell(r, c)
        cell.fill.background()
        run = cell.text_frame.paragraphs[0].add_run()
        run.text = val
        run.font.size = Pt(11 if c == 0 else 18)
        run.font.name = "IBM Plex Mono" if c == 0 else FONT
        run.font.color.rgb = MUTED if c == 0 else INK

prs.save("04-deck.pptx")
print("wrote 04-deck.pptx")
